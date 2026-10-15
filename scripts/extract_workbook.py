"""Extract numeric raw measurements, never chart aggregates or formula errors."""
import sys,json,datetime,collections
import openpyxl
w=openpyxl.load_workbook(sys.argv[1],read_only=True,data_only=True)
s=w.worksheets[0]
rows=list(s.iter_rows(values_only=True))
lines=[];items={};records=[];line=None;item=None;start=0;instrument='';invalid=0
for rn,row in enumerate(rows[:1342],1):
 b,c,d=(row[i] if len(row)>i else None for i in (1,2,3))
 if isinstance(b,str):
  compact=''.join(b.split())
  if compact=='공정검사':line='process';lines.append({'id':line,'name':'공정','type':'공정'})
  elif compact.startswith('저장탱크'):
   suffix=compact.removeprefix('저장탱크').removeprefix('TK');line='tk-'+suffix.lower();lines.append({'id':line,'name':'저장탱크 '+suffix,'type':'저장탱크'})
  elif compact=='구분':item=None
 if c:instrument=''.join(str(c).split())
 if line and isinstance(d,str) and d.strip():
  item=d.replace('\n','').strip();start=rn
  if item not in items:
   ident='item-'+str(len(items)+1)
   # Unit of HALO raw values is not explicitly specified; require confirmation.
   items[item]={'id':ident,'name':item,'unit':'℃' if item=='수분' else ('' if item=='미세수분' else 'ppm'),'instrument':instrument,'lower':None,'upper':None}
 if not line or not item:continue
 for ci,v in enumerate(row[4:376],5):
  if not isinstance(v,(int,float)) or isinstance(v,bool):continue
  month=(ci-5)//31+1;day=(ci-5)%31+1
  try:date=datetime.date(2026,month,day).isoformat()
  except ValueError:invalid+=1;continue
  records.append({'id':f'xlsx-{rn}-{ci}','lineId':line,'itemId':items[item]['id'],'date':date,'sequence':rn-start+1,'value':v,'note':'','sourceCell':f'{openpyxl.utils.get_column_letter(ci)}{rn}'})
out={'lines':lines,'items':list(items.values()),'records':records,'source':'2026년 L-CO₂ 품질검사','excludedInvalidDates':invalid}
with open(sys.argv[2],'w') as f:json.dump(out,f,ensure_ascii=False,separators=(',',':'))
print('lines',len(lines),'items',len(items),'records',len(records),'invalid-date cells',invalid)
print('records per line',dict(collections.Counter(r['lineId'] for r in records)))
